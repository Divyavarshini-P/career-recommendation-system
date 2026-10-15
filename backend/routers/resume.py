from fastapi import APIRouter, UploadFile, File, HTTPException
import pdfplumber
from docx import Document
from io import BytesIO
import os

from utils.skill_extractor import extract_skills


router = APIRouter()


# =====================================================
# Upload Resume
# =====================================================

@router.post("/upload-resume")
async def upload_resume(file: UploadFile = File(...)):

    # -------------------------------------------------
    # Check File
    # -------------------------------------------------

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No file selected."
        )

    # Get file extension

    filename = file.filename.lower()

    # -------------------------------------------------
    # Allow only PDF and DOCX
    # -------------------------------------------------

    if not (
        filename.endswith(".pdf")
        or filename.endswith(".docx")
    ):

        raise HTTPException(
            status_code=400,
            detail="Only PDF and DOCX files are supported."
        )

    try:

        # -------------------------------------------------
        # Read Uploaded File
        # -------------------------------------------------

        file_content = await file.read()

        if not file_content:

            raise HTTPException(
                status_code=400,
                detail="Uploaded file is empty."
            )

        # -------------------------------------------------
        # Extract Text
        # -------------------------------------------------

        extracted_text = ""

        # =================================================
        # PDF
        # =================================================

        if filename.endswith(".pdf"):

            with pdfplumber.open(
                BytesIO(file_content)
            ) as pdf:

                for page in pdf.pages:

                    page_text = page.extract_text()

                    if page_text:

                        extracted_text += (
                            page_text + "\n"
                        )

        # =================================================
        # DOCX
        # =================================================

        elif filename.endswith(".docx"):

            document = Document(
                BytesIO(file_content)
            )

            # Extract paragraphs

            for paragraph in document.paragraphs:

                if paragraph.text.strip():

                    extracted_text += (
                        paragraph.text + "\n"
                    )

            # Extract tables

            for table in document.tables:

                for row in table.rows:

                    row_text = []

                    for cell in row.cells:

                        row_text.append(
                            cell.text.strip()
                        )

                    extracted_text += (
                        " ".join(row_text) + "\n"
                    )

        # -------------------------------------------------
        # Clean Text
        # -------------------------------------------------

        extracted_text = extracted_text.strip()

        # -------------------------------------------------
        # Check Extraction
        # -------------------------------------------------

        if not extracted_text:

            raise HTTPException(
                status_code=400,
                detail=(
                    "Could not extract text from this resume. "
                    "Make sure the PDF is not a scanned image."
                )
            )

        # -------------------------------------------------
        # Extract Skills
        # -------------------------------------------------

        skills = extract_skills(
            extracted_text
        )

        # -------------------------------------------------
        # Return Result
        # -------------------------------------------------

        return {

            "success": True,

            "filename": file.filename,

            "extracted_text": extracted_text,

            "skills": skills,

            "word_count": len(
                extracted_text.split()
            )

        }

    except HTTPException:
        raise

    except Exception as e:

        print(
            "Resume extraction error:",
            str(e)
        )

        raise HTTPException(
            status_code=500,
            detail=(
                "Error processing resume: "
                + str(e)
            )
        )